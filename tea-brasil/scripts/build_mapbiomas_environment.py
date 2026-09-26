#!/usr/bin/env python3
"""Build 2022 municipality environmental indicators from MapBiomas Brazil Collection 11.

Official source: MAPBIOMAS_BRAZIL-COL.11-BIOME_STATE_MUNICIPALITY.zip
Year used: 2022, aligned with Censo 2022 TEA outcomes.

Definitions:
- natural vegetation = leaf classes whose level-1 group is either
  '1. Forest' or '2. Herbaceous and Shrubby Vegetation'. This deliberately
  excludes water and natural non-vegetated surfaces.
- urban area = MapBiomas class 24 ('Urban Area').
- denominator = sum of all mapped leaf-class areas in 2022 for the municipality.
Municipalities split across biomes are summed by IBGE geocode before percentages.

Territorial compatibility rule:
The Collection 11 municipal workbook uses the current municipal grid and therefore
contains Boa Esperanca do Norte (5101837), created after the 2022 Census reference
frame, plus the non-municipal RS water polygons Lagoa Mirim (4300001) and Lagoa dos
Patos (4300002). It does not contain Fernando de Noronha (2605459).
To avoid assigning post-2022 boundaries to 2022 Census outcomes, Sorriso (5107925)
and Nova Ubirata (5106240), the source municipalities affected by the creation of
Boa Esperanca do Norte, are kept in the 5,570-row output but environmental values
are NA. Fernando de Noronha is likewise kept as NA because it is absent from the
MapBiomas municipality workbook. No value is imputed.
"""
from __future__ import annotations
import csv, io, json, sys, unicodedata, urllib.parse, urllib.request, zipfile
from collections import defaultdict
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import load_workbook

EXPECTED=5570
YEAR='y2022'
FILE_ID='1otOqymHuixvkRGVl65zTTNyfaHo46Gqk'
SOURCE_URL=f'https://drive.google.com/uc?id={FILE_ID}&export=download'
SOURCE_PAGE='https://brasil.mapbiomas.org/downloads/estatisticas/'
ROOT=Path(__file__).resolve().parents[1]
BASE_JSON=ROOT/'data'/'municipios_tea_renda_2022.json'
OUT_CSV=ROOT/'data'/'mapbiomas_ambiente_2022.csv'
OUT_JSON=ROOT/'data'/'mapbiomas_ambiente_2022.json'
META=ROOT/'data'/'metadata_mapbiomas_ambiente_2022.json'
UA={'User-Agent':'TEA-Brasil/1.0 (+https://ernandes-sobreira.github.io/justa-MT/tea-brasil/)'}

# Differences expected between the Censo 2022 municipality frame and the current
# municipality workbook distributed with MapBiomas Collection 11.
EXPECTED_SOURCE_MISSING={'2605459'}  # Fernando de Noronha
EXPECTED_SOURCE_EXTRA={'4300001','4300002','5101837'}  # Lagoa Mirim, Lagoa dos Patos, Boa Esperanca do Norte
TERRITORIAL_NA={
    '2605459':('na_fonte','Fernando de Noronha não aparece na planilha municipal MapBiomas Coleção 11 usada nesta integração.'),
    '5106240':('na_incompatibilidade_territorial','Nova Ubiratã: a planilha MapBiomas atual usa limites posteriores ao Censo 2022 após a criação de Boa Esperança do Norte; valor não forçado.'),
    '5107925':('na_incompatibilidade_territorial','Sorriso: a planilha MapBiomas atual usa limites posteriores ao Censo 2022 após a criação de Boa Esperança do Norte; valor não forçado.'),
}

class DownloadForm(HTMLParser):
    def __init__(self):
        super().__init__();self.action=None;self.inputs={};self.in_form=False
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag.lower()=='form' and a.get('id')=='download-form':
            self.in_form=True;self.action=a.get('action')
        elif self.in_form and tag.lower()=='input' and a.get('name'):
            self.inputs[a['name']]=a.get('value','')
    def handle_endtag(self,tag):
        if tag.lower()=='form':self.in_form=False

def fetch(url,params=None,timeout=300):
    if params:url += ('&' if '?' in url else '?')+urllib.parse.urlencode(params)
    req=urllib.request.Request(url,headers=UA)
    with urllib.request.urlopen(req,timeout=timeout) as r:return r.read(),r.headers.get('content-type',''),r.geturl()

def download_zip():
    data,ctype,final=fetch(SOURCE_URL)
    if 'text/html' in ctype.lower() or data.lstrip().startswith(b'<!DOCTYPE html'):
        p=DownloadForm();p.feed(data.decode('utf-8',errors='replace'))
        if not p.action:raise RuntimeError('Google Drive confirmation form not found')
        data,ctype,final=fetch(p.action,p.inputs)
    if not data.startswith(b'PK'):raise RuntimeError(f'Official MapBiomas download is not ZIP/XLSX: {ctype} {len(data)} bytes')
    return data,final

def fold(v):
    return unicodedata.normalize('NFKD',str(v or '')).encode('ascii','ignore').decode('ascii').lower().strip()

def num(v):
    if v in (None,''):return 0.0
    try:return float(v)
    except:return 0.0

def main():
    base=json.loads(BASE_JSON.read_text(encoding='utf-8'))
    base_codes={str(r['codigo_ibge']) for r in base}
    if len(base)!=EXPECTED or len(base_codes)!=EXPECTED:raise RuntimeError('Reference base is not 5570 unique municipalities')

    print('Downloading official MapBiomas Collection 11 municipality statistics...',flush=True)
    raw,final_url=download_zip();print(f'Downloaded {len(raw):,} bytes',flush=True)
    z=zipfile.ZipFile(io.BytesIO(raw));books=[n for n in z.namelist() if n.lower().endswith('.xlsx')]
    if len(books)!=1:raise RuntimeError(f'Expected one XLSX, found {books}')
    wb=load_workbook(io.BytesIO(z.read(books[0])),read_only=True,data_only=True)
    if 'COVERAGE_11' not in wb.sheetnames:raise RuntimeError(f'COVERAGE_11 missing; sheets={wb.sheetnames}')
    ws=wb['COVERAGE_11'];rows=ws.iter_rows(values_only=True);header=next(rows)
    cols={str(v):i for i,v in enumerate(header)}
    required=['geocode','municipality','state','class','class_level_1',YEAR]
    for k in required:
        if k not in cols:raise RuntimeError(f'Column {k} missing')

    # The source may repeat municipality rows by biome. Aggregate by geocode.
    acc=defaultdict(lambda:{'total':0.0,'natural_veg':0.0,'urban':0.0,'biomes':set(),'name':'','state':''})
    nrows=0
    for row in rows:
        nrows+=1
        code=str(row[cols['geocode']] or '').strip()
        if not code:continue
        if code.endswith('.0'):code=code[:-2]
        code=''.join(ch for ch in code if ch.isdigit())
        if len(code)!=7:continue
        area=max(0.0,num(row[cols[YEAR]]));a=acc[code];a['total']+=area
        a['name']=str(row[cols['municipality']] or a['name']);a['state']=str(row[cols['state']] or a['state'])
        if 'biome' in cols and row[cols['biome']]:a['biomes'].add(str(row[cols['biome']]))
        level1=fold(row[cols['class_level_1']]);classid=str(row[cols['class']] or '').strip()
        if classid.endswith('.0'):classid=classid[:-2]
        if level1.startswith('1. forest') or level1.startswith('2. herbaceous and shrubby vegetation'):
            a['natural_veg']+=area
        if classid=='24':a['urban']+=area

    source_codes=set(acc)
    missing=base_codes-source_codes;extra=source_codes-base_codes
    if missing!=EXPECTED_SOURCE_MISSING or extra!=EXPECTED_SOURCE_EXTRA:
        raise RuntimeError(f'Unexpected MapBiomas/Censo code mismatch: missing={sorted(missing)} extra={sorted(extra)}')

    usable_codes=base_codes-set(TERRITORIAL_NA)
    missing_usable=usable_codes-source_codes
    if missing_usable:raise RuntimeError(f'Usable Censo 2022 codes absent from MapBiomas: {sorted(missing_usable)}')

    base_by={str(r['codigo_ibge']):r for r in base};out=[];bad_total=[]
    for code in sorted(base_codes):
        b=base_by[code]
        if code in TERRITORIAL_NA:
            status,note=TERRITORIAL_NA[code]
            a=acc.get(code,{'biomes':set()})
            out.append({
              'codigo_ibge':code,'municipio':b.get('municipio',''),'uf':b.get('uf',''),'populacao_2022':b.get('populacao_2022'),'percentual_tea_2022':b.get('percentual_tea_2022'),
              'area_mapeada_ha_2022':None,'vegetacao_natural_ha_2022':None,'vegetacao_natural_pct_2022':None,
              'area_urbanizada_ha_2022':None,'area_urbanizada_pct_2022':None,'biomas_mapbiomas':' | '.join(sorted(a.get('biomes',set()))),
              'status_ambiente_2022':status,'nota_ambiente_2022':note
            })
            continue
        a=acc[code];total=a['total']
        if total<=0:bad_total.append(code);continue
        veg_pct=a['natural_veg']/total*100;urb_pct=a['urban']/total*100
        out.append({
          'codigo_ibge':code,'municipio':b.get('municipio',''),'uf':b.get('uf',''),'populacao_2022':b.get('populacao_2022'),'percentual_tea_2022':b.get('percentual_tea_2022'),
          'area_mapeada_ha_2022':total,'vegetacao_natural_ha_2022':a['natural_veg'],'vegetacao_natural_pct_2022':veg_pct,
          'area_urbanizada_ha_2022':a['urban'],'area_urbanizada_pct_2022':urb_pct,'biomas_mapbiomas':' | '.join(sorted(a['biomes'])),
          'status_ambiente_2022':'ok','nota_ambiente_2022':''
        })
    if bad_total:raise RuntimeError(f'Municipalities with zero mapped area: {bad_total}')
    if len(out)!=EXPECTED or len({r['codigo_ibge'] for r in out})!=EXPECTED:raise RuntimeError('Output lost Censo 2022 municipalities')

    fields=list(out[0])
    with OUT_CSV.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(out)
    OUT_JSON.write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
    meta={
      'generated_at_utc':datetime.now(timezone.utc).isoformat(),'source':'MapBiomas Brasil','collection':11,'reference_year':2022,
      'territorial_frame_outcome':'Censo Demográfico 2022: 5.570 municípios','territorial_frame_mapbiomas':'Coleção 11 distribuída em 2026, com malha municipal atualizada',
      'source_page':SOURCE_PAGE,'official_download_file_id':FILE_ID,'official_download_url':SOURCE_URL,'resolved_download_url':final_url,
      'source_workbook':books[0],'source_sheet':'COVERAGE_11','source_rows_scanned':nrows,
      'expected_municipalities':EXPECTED,'municipalities_in_file':len(out),'unique_codes':len({r['codigo_ibge'] for r in out}),
      'valid_environment_municipalities':sum(1 for r in out if r['status_ambiente_2022']=='ok'),
      'missing_vegetation_cells':sum(1 for r in out if r['vegetacao_natural_pct_2022'] is None),'missing_urban_cells':sum(1 for r in out if r['area_urbanizada_pct_2022'] is None),
      'expected_source_missing_codes':sorted(EXPECTED_SOURCE_MISSING),'expected_source_extra_codes':sorted(EXPECTED_SOURCE_EXTRA),
      'intentional_na_codes':{code:{'status':TERRITORIAL_NA[code][0],'note':TERRITORIAL_NA[code][1]} for code in sorted(TERRITORIAL_NA)},
      'definitions':{
        'vegetacao_natural_pct_2022':'Sum of 2022 leaf-class hectares whose class_level_1 is Forest or Herbaceous and Shrubby Vegetation, divided by all mapped leaf-class hectares in the municipality. Water and natural non-vegetated classes are excluded.',
        'area_urbanizada_pct_2022':'MapBiomas class 24 Urban Area hectares in 2022 divided by all mapped leaf-class hectares in the municipality.',
        'municipality_cross_biome_rule':'Rows are summed by 7-digit IBGE geocode across all biome portions before percentages are calculated.',
        'territorial_compatibility_rule':'The output always preserves the 5,570 Censo 2022 codes. Values are left NA, never imputed, where the current MapBiomas municipal grid cannot be matched safely to the 2022 Census municipal frame.'
      },
      'rules':['Exactly 5570 Censo 2022 municipality codes must be present in the published file.','No municipality is imputed.','All environmental class areas use MapBiomas year 2022.','Known territorial incompatibilities are explicit NA values with notes, not zeros.']
    }
    META.write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:meta[k] for k in ['source_rows_scanned','municipalities_in_file','unique_codes','valid_environment_municipalities','missing_vegetation_cells','missing_urban_cells']},ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':
    try:main()
    except Exception as e:print('BUILD FAILED:',e,file=sys.stderr);sys.exit(1)
