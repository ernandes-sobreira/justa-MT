#!/usr/bin/env python3
"""Inspect official MapBiomas Collection 11 municipality statistics ZIP/workbooks.
No indicator is published by this probe.
"""
from __future__ import annotations
import io,json,re,urllib.parse,urllib.request,zipfile
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import load_workbook

URL='https://drive.google.com/uc?id=1otOqymHuixvkRGVl65zTTNyfaHo46Gqk&export=download'
UA={'User-Agent':'TEA-Brasil/1.0 (+https://ernandes-sobreira.github.io/justa-MT/tea-brasil/)'}

class DownloadForm(HTMLParser):
    def __init__(self):
        super().__init__();self.action=None;self.method='get';self.inputs={};self.in_form=False
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag.lower()=='form' and a.get('id')=='download-form':
            self.in_form=True;self.action=a.get('action');self.method=a.get('method','get').lower()
        elif self.in_form and tag.lower()=='input' and a.get('name'):
            self.inputs[a['name']]=a.get('value','')
    def handle_endtag(self,tag):
        if tag.lower()=='form' and self.in_form:self.in_form=False

def fetch(url,params=None):
    if params:
        url += ('&' if '?' in url else '?')+urllib.parse.urlencode(params)
    req=urllib.request.Request(url,headers=UA)
    with urllib.request.urlopen(req,timeout=300) as r:
        return r.read(),r.headers.get('content-type',''),r.geturl()

data,ctype,final=fetch(URL)
if data.startswith(b'<!DOCTYPE html') or 'text/html' in ctype.lower():
    html=data.decode('utf-8',errors='replace');p=DownloadForm();p.feed(html)
    if not p.action: raise SystemExit('Google Drive confirmation form not found')
    print('CONFIRM FORM',p.action,p.inputs)
    data,ctype,final=fetch(p.action,p.inputs)

print(json.dumps({'bytes':len(data),'content_type':ctype,'final_url':final,'head_hex':data[:16].hex()},ensure_ascii=False))
if not data.startswith(b'PK'):
    print(data[:2500].decode('utf-8',errors='replace'));raise SystemExit('confirmed download is not ZIP/XLSX')

z=zipfile.ZipFile(io.BytesIO(data))
names=z.namelist()
print('ZIP FILES',len(names))
for n in names[:200]:print('FILE',n,z.getinfo(n).file_size)

books=[n for n in names if n.lower().endswith(('.xlsx','.xlsm'))]
if not books:
    # The downloaded file could itself be an xlsx rather than an outer zip.
    try:
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        books=['<root-xlsx>']
    except Exception:
        raise SystemExit('No XLSX workbooks found in MapBiomas ZIP')
else:
    wb=None

for book in books[:20]:
    print('\n######## WORKBOOK',book)
    if book!='<root-xlsx>':
        raw=z.read(book);wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    print('SHEETS',wb.sheetnames)
    for name in wb.sheetnames:
        ws=wb[name]
        print('\n### SHEET',name,'rows',ws.max_row,'cols',ws.max_column)
        for i,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,10),values_only=True),1):
            print('ROW',i,repr(row[:25]))
